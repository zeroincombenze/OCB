__author__ = 'carlo'

import xlrd
import openpyxl
import csv

FILENAME = 'AbiCab.xlsx'
OPEN_SHEET = 'Foglio1'
HEADER_LINE = 1

FILE_COLUMNS = {
    'abi': 1,
    'cab': 2,
    'name': 3,
}

bank = {
    'id': [],
    'abi': [],
    'cab': [],
    'name': [],
}


class main():

    @staticmethod
    def add_code(abi, cab, name):
        bank['id'].append('base.' + abi + cab)
        bank['abi'].append(abi)
        bank['cab'].append(cab)
        bank['name'].append(name)
        return

    @staticmethod
    def write_file(dic, write_file):
        writer = open(write_file, 'w')
        row = ''

        for col in range(len(dic.keys())):
            row += '"' + dic.keys()[col] + '"'
            if col != len(dic.keys()) - 1:
                row += ','

        row += '\n'
        writer.write(row)
        row = ''
        for lin in range(len(dic['id'])):
            for col in range(len(dic.keys())):
                row += '"' + str(dic[dic.keys()[col]][lin]).replace('"', '\'') + '"'
                if col != len(dic.keys()) - 1:
                    row += ','
            row += '\n'
            writer.write(row)
            row = ''

        writer.close()

    @staticmethod
    def import_customers():
        import pdb;pdb.set_trace()
        wb = openpyxl.load_workbook(filename=FILENAME, read_only=True)
        ws = wb['Foglio1']
        main_bank_name = ''
        for row in range(0, ws.max_row + 1):# ws.rows:

            if row <= HEADER_LINE:
                continue
            abi = ws.cell(row, FILE_COLUMNS['abi']).value
            cab = ws.cell(row, FILE_COLUMNS['cab']).value
            name = ws.cell(row, FILE_COLUMNS['name']).value.encode('utf-8', 'ignore')

            if cab == 0:
                main_bank_name = name + ' - '
                continue

            main.add_code("{abi:05}".format(abi=abi), "{cab:05}".format(cab=cab), main_bank_name + name)

        main.write_file(bank, 'res.bank2.csv')


if __name__ == "__main__":
    main.import_customers()
    print 'END IMPORT'